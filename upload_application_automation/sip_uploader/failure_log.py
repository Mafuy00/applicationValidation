"""
Step 7's failure path: record unsuccessful uploads in `logs/Failed_upload.xlsx`
and make the failure impossible to miss in the terminal.

Rows are inserted newest-first, matching the ordering staff already expect
from the validation automation's CSV.
"""
import os
from datetime import datetime

from . import config

HEADERS = [
    "logged_at",
    "step",
    "file_name",
    "file_path",
    "sip_batch",
    "email_subject",
    "sender",
    "error_type",
    "error_message",
    "screenshot",
]

# Rendered widths, in the same order as HEADERS.
_COLUMN_WIDTHS = [20, 22, 38, 60, 14, 45, 32, 24, 70, 60]


def _require_openpyxl():
    try:
        import openpyxl  # noqa: F401
    except ImportError as e:
        raise ImportError(
            "openpyxl is required to write the Failed_upload workbook. "
            "Install it with: pip install -r requirements.txt"
        ) from e
    return openpyxl


def _new_workbook(openpyxl):
    from openpyxl.styles import Font

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Failed uploads"
    ws.append(HEADERS)
    for cell in ws[1]:
        cell.font = Font(bold=True)
    ws.freeze_panes = "A2"
    for idx, width in enumerate(_COLUMN_WIDTHS, start=1):
        ws.column_dimensions[openpyxl.utils.get_column_letter(idx)].width = width
    return wb


def _set_aside(path, reason):
    """
    Renames an unusable workbook out of the way so a fresh one can be written
    without destroying what was already there.
    """
    stamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    target = f"{path}.{reason}_{stamp}"
    try:
        os.replace(path, target)
        print(f"[warn] Existing {os.path.basename(path)} was {reason}; kept as {os.path.basename(target)}.")
    except OSError:
        pass


def _load_workbook(openpyxl, path):
    """Opens the existing workbook, falling back to a fresh one if unusable."""
    if not os.path.exists(path):
        return _new_workbook(openpyxl)
    try:
        wb = openpyxl.load_workbook(path)
    except Exception:
        # A corrupt or half-written workbook must not stop us recording the
        # failure; move it aside and start a clean one.
        _set_aside(path, "corrupt")
        return _new_workbook(openpyxl)

    ws = wb.active
    if ws.max_row < 1 or [c.value for c in ws[1]] != HEADERS:
        # Schema drifted (older file, manual edit). Writing into it would
        # misalign columns, and returning a new workbook would overwrite it on
        # save - so preserve the old file under a new name first.
        _set_aside(path, "unrecognised")
        return _new_workbook(openpyxl)
    return wb


def log_failure(
    step,
    error,
    file_path=None,
    email_subject="",
    sender="",
    screenshot_path="",
    path=None,
):
    """
    Appends one failure row to Failed_upload.xlsx (newest first) and returns
    the workbook path.

    step: which stage failed, e.g. "login" or "select_sip_batch" - this is the
        first thing staff need in order to know whether to retry or escalate.
    """
    openpyxl = _require_openpyxl()
    path = path or config.FAILED_UPLOAD_XLSX_PATH
    os.makedirs(os.path.dirname(path), exist_ok=True)

    if isinstance(error, BaseException):
        error_type = type(error).__name__
        error_message = str(error)
    else:
        error_type = ""
        error_message = str(error)

    row = [
        datetime.now().strftime("%Y-%m-%d %H:%M:%S"),
        step,
        os.path.basename(file_path) if file_path else "",
        file_path or "",
        config.SIP_BATCH_LABEL,
        email_subject or "",
        sender or "",
        error_type,
        error_message,
        screenshot_path or "",
    ]

    wb = _load_workbook(openpyxl, path)
    ws = wb.active
    ws.insert_rows(2)
    for col, value in enumerate(row, start=1):
        ws.cell(row=2, column=col, value=value)
    wb.save(path)
    return path


def notify_terminal(
    step, error, file_path=None, workbook_path=None, screenshot_path=None, dry_run=False
):
    """Prints a clearly delimited failure notice to the terminal."""
    bar = "=" * 72
    print(f"\n{bar}")
    print("  SIP UPLOAD DRY RUN FAILED (nothing was uploaded)" if dry_run else "  SIP UPLOAD FAILED")
    print(bar)
    print(f"  Step        : {step}")
    if file_path:
        print(f"  Form        : {file_path}")
    prefix = f"{type(error).__name__}: " if isinstance(error, BaseException) else ""
    print(f"  Error       : {prefix}{error}")
    if screenshot_path:
        print(f"  Screenshot  : {screenshot_path}")
    if workbook_path:
        print(f"  Logged to   : {workbook_path}")
    print(f"{bar}\n")


def record_failure(
    step,
    error,
    file_path=None,
    email_subject="",
    sender="",
    screenshot_path="",
    dry_run=False,
):
    """
    Logs the failure to the workbook and notifies the terminal.

    dry_run: write to the separate dry-run workbook, so rehearsals never
        pollute the Failed_upload.xlsx staff actually triage from.
    """
    workbook_path = None
    try:
        workbook_path = log_failure(
            step=step,
            error=error,
            file_path=file_path,
            email_subject=email_subject,
            sender=sender,
            screenshot_path=screenshot_path,
            path=config.FAILED_UPLOAD_DRYRUN_XLSX_PATH if dry_run else None,
        )
    except Exception as log_error:
        # Never let a logging problem swallow the original failure.
        print(f"[warn] Could not write Failed_upload workbook: {log_error}")

    notify_terminal(
        step=step,
        error=error,
        file_path=file_path,
        workbook_path=workbook_path,
        screenshot_path=screenshot_path,
        dry_run=dry_run,
    )
    return workbook_path
